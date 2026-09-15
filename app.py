import io
import os
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from flask import Flask, jsonify, render_template, request, send_file
from flask_sqlalchemy import SQLAlchemy
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import or_

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DATA_DIR.mkdir(exist_ok=True)
SEED_FILE = DATA_DIR / "RA_seed.xlsx"

app = Flask(__name__)
app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", "jt-ra-dashboard-dev")

database_url = os.getenv("DATABASE_URL")
if not database_url:
    database_url = f"sqlite:///{(DATA_DIR / 'dashboard.db').as_posix()}"
elif database_url.startswith("postgres://"):
    database_url = database_url.replace("postgres://", "postgresql+psycopg://", 1)
elif database_url.startswith("postgresql://") and "+psycopg" not in database_url:
    database_url = database_url.replace("postgresql://", "postgresql+psycopg://", 1)

app.config["SQLALCHEMY_DATABASE_URI"] = database_url
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["SQLALCHEMY_ENGINE_OPTIONS"] = {"pool_pre_ping": True}
db = SQLAlchemy(app)

SAO_PAULO = ZoneInfo("America/Sao_Paulo")


def now_sp_naive():
    return datetime.now(SAO_PAULO).replace(tzinfo=None)


class Complaint(db.Model):
    __tablename__ = "complaints"
    id = db.Column(db.Integer, primary_key=True)
    ra_id = db.Column(db.String(80), index=True)
    complaint_date = db.Column(db.DateTime, index=True)
    regional = db.Column(db.String(30), index=True)
    owner = db.Column(db.String(160))
    tracking = db.Column(db.String(100))
    driver = db.Column(db.String(220))
    base = db.Column(db.String(160), index=True)
    base_type = db.Column(db.String(30), index=True)
    rm = db.Column(db.String(160), index=True)
    directed_date = db.Column(db.DateTime)
    reason = db.Column(db.String(220), index=True)
    solution = db.Column(db.String(220), index=True)
    deadline_manual = db.Column(db.String(40))
    treatment = db.Column(db.Text)
    observation = db.Column(db.Text)
    source_row = db.Column(db.Integer)
    updated_at = db.Column(db.DateTime, default=now_sp_naive, onupdate=now_sp_naive)


class AppMeta(db.Model):
    __tablename__ = "app_meta"
    key = db.Column(db.String(80), primary_key=True)
    value = db.Column(db.String(500), nullable=False)


def normalize_text(value):
    if value is None:
        return ""
    return str(value).strip()


def normalize_key(value):
    text = normalize_text(value).upper()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^A-Z0-9]+", " ", text).strip()


def classify_base(base):
    b = normalize_text(base).upper()
    return "FRANQUIA" if re.match(r"^F\s+", b) else "BASE PRÓPRIA"


def normalize_reason(value):
    raw = normalize_text(value)
    key = normalize_key(raw)
    aliases = {
        "POSTURA DO MOTORISTA": "Postura do Motorista",
        "POSTURA": "Postura",
        "TROCA DE ETIQUETA": "Troca de etiqueta",
        "ATRASO NA ENTREGA": "Atraso na entrega",
        "ACAREACAO": "Acareação",
        "PONTO DE COLETA": "Ponto de Coleta",
        "INFORMACAO": "Informação",
        "AVARIA": "Avaria",
        "EXTRAVIO AVARIA": "Extravio/Avaria",
        "ITEM FALTANTE": "Item faltante",
        "EM ROTA": "Em rota",
        "TRABALHISTA": "Trabalhista",
    }
    return aliases.get(key, raw)


TREATMENT_BY_REASON = {
    "Atraso na entrega": "Validar última movimentação, acionar a base responsável e registrar previsão de entrega.",
    "Acareação": "Realizar acareação com base/motorista, registrar evidências e concluir a devolutiva ao cliente.",
    "Postura": "Apurar a conduta, identificar o responsável e registrar orientação/correção aplicada.",
    "Postura do Motorista": "Apurar a conduta do motorista, registrar evidências e aplicar orientação/correção.",
    "Trabalhista": "Direcionar para a área responsável e registrar o protocolo da tratativa.",
    "Ponto de Coleta": "Validar o ponto de coleta, responsável local e evidências do atendimento.",
    "Informação": "Confirmar a informação correta, ajustar o registro e retornar ao cliente.",
    "Avaria": "Validar evidências da avaria, responsabilidade operacional e fluxo de ressarcimento quando aplicável.",
    "Extravio/Avaria": "Realizar busca operacional, validar evidências e direcionar o fluxo de extravio/avaria.",
    "Troca de etiqueta": "Rastrear as etiquetas envolvidas, corrigir a vinculação e validar o destino correto.",
    "Item faltante": "Conferir quantidade, peso, volumetria e evidências para localizar o item faltante.",
    "Em rota": "Validar rota, motorista e previsão de conclusão da entrega.",
}
DEFAULT_TREATMENT = "Analisar o motivo, acionar o responsável e registrar evidências e devolutiva da tratativa."


def suggested_treatment(reason):
    return TREATMENT_BY_REASON.get(normalize_reason(reason), DEFAULT_TREATMENT)


def parse_datetime(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if hasattr(value, "year") and hasattr(value, "month") and hasattr(value, "day"):
        try:
            return datetime(value.year, value.month, value.day)
        except Exception:
            pass
    text = normalize_text(value)
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def format_dt(value):
    return value.strftime("%Y-%m-%dT%H:%M") if value else ""


def format_date_br(value):
    return value.strftime("%d/%m/%Y %H:%M") if value else ""


def sla_auto(complaint_date):
    if not complaint_date:
        return "SEM DATA"
    delta = now_sp_naive() - complaint_date
    return "DENTRO DO PRAZO" if delta <= timedelta(hours=24) else "FORA DO PRAZO"


def is_open(solution):
    value = normalize_key(solution)
    return not value or value == "EM ABERTO"


def get_version():
    meta = db.session.get(AppMeta, "data_version")
    return int(meta.value) if meta else 0


def bump_version():
    meta = db.session.get(AppMeta, "data_version")
    if not meta:
        meta = AppMeta(key="data_version", value="1")
        db.session.add(meta)
    else:
        meta.value = str(int(meta.value) + 1)
    updated = db.session.get(AppMeta, "last_update")
    stamp = now_sp_naive().isoformat(timespec="seconds")
    if not updated:
        db.session.add(AppMeta(key="last_update", value=stamp))
    else:
        updated.value = stamp


def header_index(headers):
    out = {}
    for idx, h in enumerate(headers):
        key = normalize_key(h)
        if key:
            out[key] = idx
    return out


def pick(row, indexes, *names, fallback=None):
    for name in names:
        idx = indexes.get(normalize_key(name))
        if idx is not None and idx < len(row):
            return row[idx]
    if fallback is not None and fallback < len(row):
        return row[fallback]
    return None


def workbook_to_records(file_obj, preserve_existing=True):
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = list(next(rows))
    indexes = header_index(headers)

    preserved = {}
    if preserve_existing:
        for item in Complaint.query.all():
            if item.ra_id:
                preserved[item.ra_id] = {
                    "treatment": item.treatment,
                    "deadline_manual": item.deadline_manual,
                }

    records = []
    for excel_row, row in enumerate(rows, start=2):
        if not any(v not in (None, "") for v in row):
            continue
        ra_raw = pick(row, indexes, "ID RA", fallback=0)
        complaint_date = parse_datetime(pick(row, indexes, "DATA RECLAMAÇÃO", "DATA RECLAMACAO", fallback=1))
        regional = normalize_text(pick(row, indexes, "ESTADO", "REGIONAL", fallback=2)).upper()
        owner = normalize_text(pick(row, indexes, "RESPONSÁVEL", "RESPONSAVEL", "ATENDENTE", fallback=3))
        tracking = normalize_text(pick(row, indexes, "RASTREIO", fallback=4))
        driver = normalize_text(pick(row, indexes, "MOTORISTA", fallback=5))
        base = normalize_text(pick(row, indexes, "BASE", fallback=6))
        rm = normalize_text(pick(row, indexes, "RM", fallback=8))
        directed_date = parse_datetime(pick(row, indexes, "DATA DIRECIONAMENTO", fallback=9))
        reason = normalize_reason(pick(row, indexes, "MOTIVO", fallback=10))
        solution = normalize_text(pick(row, indexes, "SOLUÇÃO", "SOLUCAO", fallback=11))
        deadline_manual = normalize_text(pick(row, indexes, "DENTRO DO PRAZO", "PRAZO MANUAL", fallback=12))
        observation = normalize_text(pick(row, indexes, "OBSERVAÇÃO", "OBSERVACAO", fallback=13))
        treatment_raw = normalize_text(pick(row, indexes, "TRATATIVA"))
        ra_id = normalize_text(ra_raw)
        if ra_id.endswith(".0"):
            ra_id = ra_id[:-2]
        old = preserved.get(ra_id, {})
        treatment = treatment_raw or old.get("treatment") or suggested_treatment(reason)
        if not deadline_manual:
            deadline_manual = old.get("deadline_manual", "")

        records.append(Complaint(
            ra_id=ra_id,
            complaint_date=complaint_date,
            regional=regional,
            owner=owner,
            tracking=tracking,
            driver=driver,
            base=base,
            base_type=classify_base(base),
            rm=rm,
            directed_date=directed_date,
            reason=reason,
            solution=solution,
            deadline_manual=deadline_manual,
            treatment=treatment,
            observation=observation,
            source_row=excel_row,
        ))
    return records


def replace_database_from_workbook(file_obj, preserve_existing=True):
    records = workbook_to_records(file_obj, preserve_existing=preserve_existing)
    if not records:
        raise ValueError("Nenhuma linha de dados foi encontrada na planilha.")
    db.session.query(Complaint).delete()
    db.session.add_all(records)
    bump_version()
    db.session.commit()
    return len(records)


def serialize_item(item):
    return {
        "id": item.id,
        "ra_id": item.ra_id or "",
        "complaint_date": format_dt(item.complaint_date),
        "complaint_date_display": format_date_br(item.complaint_date),
        "regional": item.regional or "",
        "owner": item.owner or "",
        "tracking": item.tracking or "",
        "driver": item.driver or "",
        "base": item.base or "",
        "base_type": classify_base(item.base),
        "rm": item.rm or "",
        "directed_date": format_dt(item.directed_date),
        "reason": item.reason or "",
        "solution": item.solution or "",
        "deadline_manual": item.deadline_manual or "",
        "sla_auto": sla_auto(item.complaint_date),
        "treatment": item.treatment or suggested_treatment(item.reason),
        "observation": item.observation or "",
        "open": is_open(item.solution),
        "updated_at": format_date_br(item.updated_at),
    }


def apply_dashboard_filters(query):
    regional = normalize_text(request.args.get("regional", "TODOS")).upper()
    rm = normalize_text(request.args.get("rm", ""))
    start = parse_datetime(request.args.get("start"))
    end = parse_datetime(request.args.get("end"))
    if regional and regional != "TODOS":
        query = query.filter(Complaint.regional == regional)
    if rm:
        if rm == "__SEM_RM__":
            query = query.filter(or_(Complaint.rm.is_(None), Complaint.rm == ""))
        else:
            query = query.filter(Complaint.rm == rm)
    if start:
        query = query.filter(Complaint.complaint_date >= start)
    if end:
        query = query.filter(Complaint.complaint_date < end + timedelta(days=1))
    return query


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/health")
def health():
    return {"status": "ok", "rows": Complaint.query.count(), "version": get_version()}


@app.get("/api/version")
def api_version():
    meta = db.session.get(AppMeta, "last_update")
    return jsonify({"version": get_version(), "last_update": meta.value if meta else None})


@app.get("/api/filters")
def api_filters():
    rms = [r[0] for r in db.session.query(Complaint.rm).distinct().order_by(Complaint.rm).all() if r[0]]
    regions = [r[0] for r in db.session.query(Complaint.regional).distinct().order_by(Complaint.regional).all() if r[0]]
    reasons = [r[0] for r in db.session.query(Complaint.reason).distinct().order_by(Complaint.reason).all() if r[0]]
    solutions = [r[0] for r in db.session.query(Complaint.solution).distinct().order_by(Complaint.solution).all() if r[0]]
    has_blank_rm = Complaint.query.filter(or_(Complaint.rm.is_(None), Complaint.rm == "")).first() is not None
    return jsonify({"rms": rms, "regions": regions, "reasons": reasons, "solutions": solutions, "has_blank_rm": has_blank_rm})


@app.get("/api/dashboard")
def api_dashboard():
    items = apply_dashboard_filters(Complaint.query).all()
    total = len(items)
    open_items = [x for x in items if is_open(x.solution)]
    franchises = [x for x in items if classify_base(x.base) == "FRANQUIA"]
    own = total - len(franchises)
    within = sum(1 for x in items if sla_auto(x.complaint_date) == "DENTRO DO PRAZO")
    overdue = sum(1 for x in items if sla_auto(x.complaint_date) == "FORA DO PRAZO")
    no_date = total - within - overdue

    reason_counts = Counter(normalize_reason(x.reason) or "Sem motivo" for x in items)
    base_counts = Counter(classify_base(x.base) for x in items)

    open_by_rm = defaultdict(lambda: {"total": 0, "own": 0, "franchise": 0, "within": 0, "overdue": 0})
    for x in open_items:
        key = x.rm or "SEM RM"
        row = open_by_rm[key]
        row["total"] += 1
        if classify_base(x.base) == "FRANQUIA":
            row["franchise"] += 1
        else:
            row["own"] += 1
        sla = sla_auto(x.complaint_date)
        if sla == "DENTRO DO PRAZO":
            row["within"] += 1
        elif sla == "FORA DO PRAZO":
            row["overdue"] += 1

    open_rows = [{"rm": k, **v} for k, v in open_by_rm.items()]
    open_rows.sort(key=lambda x: (-x["total"], x["rm"]))

    recent = sorted(items, key=lambda x: x.complaint_date or datetime.min, reverse=True)[:12]

    return jsonify({
        "version": get_version(),
        "kpis": {
            "total": total,
            "open": len(open_items),
            "own": own,
            "franchise": len(franchises),
            "within": within,
            "overdue": overdue,
            "no_date": no_date,
        },
        "reason_counts": [{"label": k, "value": v} for k, v in reason_counts.most_common()],
        "base_counts": [{"label": k, "value": v} for k, v in base_counts.items()],
        "open_by_rm": open_rows,
        "recent": [serialize_item(x) for x in recent],
    })


@app.get("/api/ranking")
def api_ranking():
    only_open = request.args.get("only_open", "0") == "1"
    items = apply_dashboard_filters(Complaint.query).all()
    if only_open:
        items = [x for x in items if is_open(x.solution)]
    reasons = sorted({normalize_reason(x.reason) or "Sem motivo" for x in items})
    matrix = defaultdict(lambda: {
        "total": 0, "within": 0, "overdue": 0, "own": 0, "franchise": 0,
        "reasons": Counter()
    })
    for x in items:
        rm = x.rm or "SEM RM"
        row = matrix[rm]
        row["total"] += 1
        row["reasons"][normalize_reason(x.reason) or "Sem motivo"] += 1
        if classify_base(x.base) == "FRANQUIA":
            row["franchise"] += 1
        else:
            row["own"] += 1
        s = sla_auto(x.complaint_date)
        if s == "DENTRO DO PRAZO": row["within"] += 1
        if s == "FORA DO PRAZO": row["overdue"] += 1

    rows = []
    for rm, data in matrix.items():
        rows.append({
            "rm": rm,
            "total": data["total"],
            "within": data["within"],
            "overdue": data["overdue"],
            "own": data["own"],
            "franchise": data["franchise"],
            "reasons": {r: data["reasons"].get(r, 0) for r in reasons},
        })
    rows.sort(key=lambda x: (-x["total"], x["rm"]))
    return jsonify({"version": get_version(), "reasons": reasons, "rows": rows})


@app.get("/api/rows")
def api_rows():
    page = max(1, int(request.args.get("page", 1)))
    per_page = min(100, max(10, int(request.args.get("per_page", 30))))
    q = Complaint.query
    search = normalize_text(request.args.get("q"))
    regional = normalize_text(request.args.get("regional", "TODOS")).upper()
    rm = normalize_text(request.args.get("rm"))
    base_type = normalize_text(request.args.get("base_type"))
    if regional and regional != "TODOS":
        q = q.filter(Complaint.regional == regional)
    if rm:
        if rm == "__SEM_RM__": q = q.filter(or_(Complaint.rm.is_(None), Complaint.rm == ""))
        else: q = q.filter(Complaint.rm == rm)
    if base_type:
        q = q.filter(Complaint.base_type == base_type)
    if search:
        pattern = f"%{search}%"
        q = q.filter(or_(
            Complaint.ra_id.ilike(pattern), Complaint.tracking.ilike(pattern), Complaint.driver.ilike(pattern),
            Complaint.base.ilike(pattern), Complaint.rm.ilike(pattern), Complaint.reason.ilike(pattern),
            Complaint.solution.ilike(pattern), Complaint.observation.ilike(pattern), Complaint.treatment.ilike(pattern)
        ))
    q = q.order_by(Complaint.complaint_date.desc().nullslast(), Complaint.id.desc())
    pagination = q.paginate(page=page, per_page=per_page, error_out=False)
    return jsonify({
        "version": get_version(),
        "page": page,
        "per_page": per_page,
        "pages": pagination.pages,
        "total": pagination.total,
        "rows": [serialize_item(x) for x in pagination.items],
    })


EDITABLE_FIELDS = {
    "ra_id", "complaint_date", "regional", "owner", "tracking", "driver", "base", "rm",
    "directed_date", "reason", "solution", "deadline_manual", "treatment", "observation"
}


@app.post("/api/rows")
def create_row():
    payload = request.get_json(force=True) or {}
    item = Complaint()
    for field in EDITABLE_FIELDS:
        if field in payload:
            value = payload[field]
            if field in {"complaint_date", "directed_date"}:
                value = parse_datetime(value)
            elif field == "reason":
                value = normalize_reason(value)
            else:
                value = normalize_text(value)
            setattr(item, field, value)
    item.base_type = classify_base(item.base)
    if not item.treatment:
        item.treatment = suggested_treatment(item.reason)
    item.updated_at = now_sp_naive()
    db.session.add(item)
    bump_version()
    db.session.commit()
    return jsonify({"ok": True, "version": get_version(), "row": serialize_item(item)}), 201


@app.put("/api/rows/<int:item_id>")
def update_row(item_id):
    item = db.session.get(Complaint, item_id)
    if not item:
        return jsonify({"error": "Registro não encontrado."}), 404
    payload = request.get_json(force=True) or {}
    old_reason = item.reason
    for field in EDITABLE_FIELDS:
        if field in payload:
            value = payload[field]
            if field in {"complaint_date", "directed_date"}:
                value = parse_datetime(value)
            elif field == "reason":
                value = normalize_reason(value)
            else:
                value = normalize_text(value)
            setattr(item, field, value)
    item.base_type = classify_base(item.base)
    if "reason" in payload and normalize_reason(old_reason) != normalize_reason(item.reason) and not normalize_text(payload.get("treatment")):
        item.treatment = suggested_treatment(item.reason)
    if not item.treatment:
        item.treatment = suggested_treatment(item.reason)
    item.updated_at = now_sp_naive()
    bump_version()
    db.session.commit()
    return jsonify({"ok": True, "version": get_version(), "row": serialize_item(item)})


@app.delete("/api/rows/<int:item_id>")
def delete_row(item_id):
    item = db.session.get(Complaint, item_id)
    if not item:
        return jsonify({"error": "Registro não encontrado."}), 404
    db.session.delete(item)
    bump_version()
    db.session.commit()
    return jsonify({"ok": True, "version": get_version()})


@app.post("/api/upload")
def upload_xlsx():
    if "file" not in request.files:
        return jsonify({"error": "Selecione um arquivo XLSX."}), 400
    f = request.files["file"]
    if not f.filename.lower().endswith(".xlsx"):
        return jsonify({"error": "O arquivo precisa ser .xlsx"}), 400
    try:
        count = replace_database_from_workbook(f.stream, preserve_existing=True)
        return jsonify({"ok": True, "rows": count, "version": get_version()})
    except Exception as exc:
        db.session.rollback()
        return jsonify({"error": f"Não foi possível importar a planilha: {exc}"}), 400


@app.get("/api/export")
def export_xlsx():
    items = Complaint.query.order_by(Complaint.id).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Reclame Aqui"
    headers = [
        "ID RA", "DATA RECLAMAÇÃO", "ESTADO", "RESPONSÁVEL", "RASTREIO", "Motorista", "BASE",
        "BASE/FRANQUIA", "RM", "DATA DIRECIONAMENTO", "Motivo", "SOLUÇÃO", "DENTRO DO PRAZO",
        "TRATATIVA", "Observação", "SLA 24H AUTOMÁTICO"
    ]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="D71920")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for x in items:
        ws.append([
            x.ra_id, x.complaint_date, x.regional, x.owner, x.tracking, x.driver, x.base,
            classify_base(x.base), x.rm, x.directed_date, x.reason, x.solution, x.deadline_manual,
            x.treatment, x.observation, sla_auto(x.complaint_date)
        ])
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "dd/mm/yyyy hh:mm"
        row[9].number_format = "dd/mm/yyyy hh:mm"
        row[7].font = Font(color="00853F" if row[7].value == "FRANQUIA" else "D71920")
    widths = [14, 20, 10, 18, 22, 28, 18, 18, 22, 20, 24, 20, 20, 58, 48, 22]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, as_attachment=True, download_name="reclame_aqui_dashboard.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/api/treatment-suggestion")
def treatment_suggestion():
    reason = request.args.get("reason", "")
    return jsonify({"reason": normalize_reason(reason), "treatment": suggested_treatment(reason)})


def bootstrap():
    db.create_all()
    if not db.session.get(AppMeta, "data_version"):
        db.session.add(AppMeta(key="data_version", value="0"))
        db.session.commit()
    if Complaint.query.count() == 0 and SEED_FILE.exists():
        try:
            with SEED_FILE.open("rb") as f:
                replace_database_from_workbook(f, preserve_existing=False)
            print(f"Seed carregado: {Complaint.query.count()} registros.")
        except Exception as exc:
            db.session.rollback()
            print(f"Falha ao carregar seed: {exc}")


with app.app_context():
    bootstrap()


if __name__ == "__main__":
    port = int(os.getenv("PORT", "5000"))
    app.run(host="0.0.0.0", port=port, debug=os.getenv("FLASK_DEBUG") == "1")
